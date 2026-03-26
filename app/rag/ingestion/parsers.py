from __future__ import annotations

import logging
from pathlib import Path
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from docling.document_converter import DocumentConverter

from app.rag.models import Document

logger = logging.getLogger(__name__)

_converter: "DocumentConverter | None" = None


def _get_converter() -> "DocumentConverter":
    global _converter
    if _converter is None:
        logger.info("Loading DocumentConverter (first run downloads models — may take 1-2 min)...")
        try:
            from docling.document_converter import DocumentConverter  # noqa: PLC0415
        except ModuleNotFoundError as exc:
            raise RuntimeError(
                "PDF/advanced document parsing requires optional dependency 'docling'. "
                "Install with: uv sync --extra pdf"
            ) from exc
        _converter = DocumentConverter()
        logger.info("DocumentConverter ready.")
    return _converter


def parse_document(path: str | Path) -> tuple[Document, str]:
    """Parse a document into a Document metadata object and structured Markdown.

    Plain-text formats (.md, .txt) are read directly.
    DOCX is parsed via python-docx.
    PDF and other advanced formats are parsed via docling (optional extra).

    Returns:
        (document, markdown_text)
    """
    path = Path(path)
    logger.info(f"Starting document parse: {path}")

    suffix = path.suffix.lower()

    if suffix in {".md", ".txt", ".markdown"}:
        logger.info(f"Plain-text file detected, reading directly (no docling needed)")
        markdown = path.read_text(encoding="utf-8")
        document = Document(
            source_file=path.name,
            doc_type=suffix.lstrip("."),
            page_count=None,
        )
        logger.info(f"Document parsed: {path.name}")
        return document, markdown

    if suffix == ".docx":
        logger.info("DOCX file detected, using lightweight parser (no docling needed)")
        try:
            from docx import Document as DocxDocument  # noqa: PLC0415
        except Exception as exc:  # pragma: no cover - depends on local runtime
            raise RuntimeError(
                "python-docx is unavailable. Install dependencies with 'uv sync'."
            ) from exc

        docx_doc = DocxDocument(str(path))
        paragraphs = [p.text.strip() for p in docx_doc.paragraphs if p.text.strip()]
        markdown = "\n\n".join(paragraphs)
        document = Document(
            source_file=path.name,
            doc_type="docx",
            page_count=None,
        )
        logger.info(f"Document parsed: {path.name} ({len(paragraphs)} paragraphs)")
        return document, markdown

    if suffix == ".epub":
        logger.info("EPUB file detected, using lightweight parser")
        try:
            import ebooklib  # noqa: PLC0415
            from bs4 import BeautifulSoup  # noqa: PLC0415
            from ebooklib import epub  # noqa: PLC0415
        except ModuleNotFoundError as exc:
            raise RuntimeError(
                "EPUB parsing requires optional dependencies. Install with: uv sync --extra epub"
            ) from exc

        book = epub.read_epub(str(path))
        sections: list[str] = []
        for item in book.get_items():
            if item.get_type() != ebooklib.ITEM_DOCUMENT:
                continue
            soup = BeautifulSoup(item.get_content(), "html.parser")
            text = soup.get_text("\n", strip=True)
            if text:
                sections.append(text)

        markdown = "\n\n".join(sections)
        document = Document(
            source_file=path.name,
            doc_type="epub",
            page_count=None,
        )
        logger.info(f"Document parsed: {path.name} ({len(sections)} sections)")
        return document, markdown

    converter = _get_converter()
    logger.info(f"Converting document via docling: {path}")
    result = converter.convert(str(path))
    doc = result.document
    logger.info(f"Document conversion complete")

    logger.debug("Exporting to Markdown...")
    markdown = doc.export_to_markdown()
    logger.debug(f"Markdown export complete (size: {len(markdown)} chars)")
    
    page_count = len(doc.pages) if hasattr(doc, "pages") and doc.pages else None

    document = Document(
        source_file=path.name,
        doc_type=path.suffix.lstrip(".").lower(),
        page_count=page_count,
    )
    logger.info(f"Document parsed: {path.name} ({page_count} pages)")
    return document, markdown
